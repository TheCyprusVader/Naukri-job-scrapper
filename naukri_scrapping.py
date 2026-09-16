import re
from datetime import date, datetime
from playwright.sync_api import sync_playwright
from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook

FILENAME = "naukri_jobs_weekly.xlsx"
HEADERS = ["title", "company", "experience", "location", "posted", "skills", "description", "link", "date_added"]

search_terms = ["data analyst", "business analyst", "MIS executive"]
num_pages = 10

keywords = ["data analyst", "business analyst", "mis", "data science", "bi analyst", "business intelligence"]


def is_fresh(posted_text):
    text = posted_text.lower()
    if "just now" in text or "hour" in text:
        return True
    match = re.search(r"(\d+)\s*day", text)
    if match:
        return int(match.group(1)) <= 1
    return False


def experience_fits(exp_text, min_exp=0, max_exp=3, max_ceiling=5):
    match = re.match(r"(\d+)-(\d+)", exp_text)
    if not match:
        return False
    job_min, job_max = int(match.group(1)), int(match.group(2))
    overlaps = job_min <= max_exp and job_max >= min_exp
    return overlaps and job_max <= max_ceiling


def load_or_create_workbook():
    try:
        wb = load_workbook(FILENAME)
    except FileNotFoundError:
        wb = Workbook()
        wb.remove(wb.active)  # remove default empty "Sheet"
    return wb


def get_all_existing_links(wb):
    """Collect every job link already present, across every sheet — used to prevent re-adding duplicates."""
    links = set()
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and len(row) >= 8 and row[7]:
                links.add(row[7])
    return links


def get_or_create_sheet(wb, name):
    if name in wb.sheetnames:
        return wb[name]
    ws = wb.create_sheet(name)
    ws.append(HEADERS)
    return ws


def combine_week(wb):
    """Merge Monday–Friday sheets into one deduped weekly summary sheet."""
    weekdays = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
    combined_name = f"Week Summary {date.today()}"

    if combined_name in wb.sheetnames:
        wb.remove(wb[combined_name])
    combined = wb.create_sheet(combined_name)
    combined.append(HEADERS)

    seen = set()
    for day in weekdays:
        if day in wb.sheetnames:
            ws = wb[day]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row and row[7] and row[7] not in seen:
                    seen.add(row[7])
                    combined.append(row)

    print(f"Weekly summary built: '{combined_name}' ({len(seen)} unique jobs)")


def scrape_jobs(existing_links):
    jobs = []
    seen_links = set(existing_links)  # start knowing what's already saved anywhere in the workbook

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=False)  # headless=True gets blocked by Naukri's bot detection (returns 0 results silently)
        page = browser.new_page()

        for term in search_terms:
            term_encoded = term.replace(" ", "%20")
            slug = term.replace(" ", "-").lower()
            base_url = f"https://www.naukri.com/{slug}-jobs"

            print(f"\n=== Searching: {term} ===")

            for page_num in range(1, num_pages + 1):
                url = f"{base_url}?k={term_encoded}" if page_num == 1 else f"{base_url}-{page_num}?k={term_encoded}"
                page.goto(url)
                page.wait_for_timeout(5000)
                html = page.content()

                soup = BeautifulSoup(html, "html.parser")
                job_cards = soup.find_all("div", class_="srp-jobtuple-wrapper")
                print(f"  Page {page_num}: {len(job_cards)} cards")

                for card in job_cards:
                    title_tag = card.find("a", class_="title")
                    title = title_tag.get_text(strip=True) if title_tag else "N/A"

                    if not any(kw in title.lower() for kw in keywords):
                        continue

                    exp_tag = card.find("span", class_="expwdth")
                    experience = exp_tag.get_text(strip=True) if exp_tag else "N/A"
                    if not experience_fits(experience):
                        continue

                    posted_tag = card.find("span", class_="job-post-day")
                    posted = posted_tag.get_text(strip=True) if posted_tag else ""
                    if not is_fresh(posted):
                        continue

                    raw_link = title_tag["href"] if title_tag else ""
                    if not raw_link or raw_link in seen_links:
                        continue
                    seen_links.add(raw_link)

                    company_tag = card.find("a", class_="comp-name")
                    loc_tag = card.find("span", class_="locWdth")
                    desc_tag = card.find("span", class_="job-desc")
                    skill_tags = card.find_all("li", class_="dot-gt")

                    jobs.append([
                        title,
                        company_tag.get_text(strip=True) if company_tag else "N/A",
                        experience,
                        loc_tag.get_text(strip=True) if loc_tag else "N/A",
                        posted,
                        ", ".join(s.get_text(strip=True) for s in skill_tags),
                        desc_tag.get_text(strip=True) if desc_tag else "N/A",
                        raw_link,
                        str(date.today()),
                    ])

        browser.close()

    return jobs


def main():
    wb = load_or_create_workbook()
    today_name = datetime.today().strftime("%A")

    if today_name == "Saturday":
        combine_week(wb)
        wb.save(FILENAME)
        print("Saturday run: weekly combine complete, no new scrape today.")
        return

    existing_links = get_all_existing_links(wb)
    jobs = scrape_jobs(existing_links)

    ws = get_or_create_sheet(wb, today_name)
    for row in jobs:
        ws.append(row)

    wb.save(FILENAME)
    print(f"\n{len(jobs)} new jobs added to '{today_name}' sheet.")
    print(f"Workbook saved as {FILENAME}")


if __name__ == "__main__":
    main()